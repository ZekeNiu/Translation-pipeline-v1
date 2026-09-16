"""DOCX generator for translated MinerU markdown."""
from __future__ import annotations

import re
import sys
from pathlib import Path

from inline_semantics import iter_styled_runs, normalize_inline_output
from table_utils import parse_html_tables
from document_blocks import REFERENCE

try:
    from docx import Document
    from docx.shared import Pt, Inches, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.enum.section import WD_SECTION_START, WD_ORIENT
except ImportError:
    import subprocess

    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    from docx import Document
    from docx.shared import Pt, Inches, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.oxml.ns import qn


BODY_SIZE = 11
HEADING_SIZES = {1: 17, 2: 14, 3: 13, 4: 12}
LINE_SPACING = 1.5
FIRST_INDENT = Cm(0.74)
CITATION_SIZE = 8
TABLE_SIZE = 8.5
CITE_RE = re.compile(r"(\[\d+(?:[,\s\u2013–-]+\d+)*\])")
XML_BAD_CHAR_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def set_font(run, bold=False, italic=False, size=None):
    run.bold = bold
    run.italic = italic
    run.font.name = "Times New Roman"
    if size:
        run.font.size = Pt(size)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = run._element.makeelement(qn("w:rFonts"), {})
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:eastAsia"), "宋体")
    rFonts.set(qn("w:ascii"), "Times New Roman")
    rFonts.set(qn("w:hAnsi"), "Times New Roman")


def _add_text_runs(paragraph, text: str, size=BODY_SIZE, bold=False, italic=False):
    text = XML_BAD_CHAR_RE.sub("", text)
    text = normalize_inline_output(text)
    for styled in iter_styled_runs(text):
        parts = CITE_RE.split(styled.text)
        for part in parts:
            if not part:
                continue
            run = paragraph.add_run(part)
            is_citation = bool(CITE_RE.fullmatch(part))
            run.font.superscript = styled.superscript or is_citation
            run.font.subscript = styled.subscript and not is_citation
            set_font(
                run,
                bold=bold or styled.bold,
                italic=italic or styled.italic,
                size=CITATION_SIZE if is_citation else size,
            )


def _paragraph(doc, text: str, size=BODY_SIZE, first_indent=True, italic=False, bold=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.first_line_indent = FIRST_INDENT if first_indent else Cm(0)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = LINE_SPACING
    pf.space_after = Pt(2)
    pf.widow_control = True
    _add_text_runs(p, text, size=size, bold=bold, italic=italic)
    return p


def _is_reference_heading(text: str) -> bool:
    return bool(REFERENCE.fullmatch(text.strip()))


def _reference_paragraph(doc, text: str):
    p = _paragraph(doc, text, size=8, first_indent=False)
    pf = p.paragraph_format
    pf.left_indent = Cm(0.45)
    pf.first_line_indent = Cm(-0.45)
    pf.space_after = Pt(1)
    pf.line_spacing = 1.0
    return p


def _safe_int(value, default=1):
    try:
        return max(default, int(value))
    except (TypeError, ValueError):
        return default


def _add_html_table(doc, html: str) -> bool:
    parsed_tables = parse_html_tables(html)
    if not parsed_tables:
        return False

    for parsed in parsed_tables:
        rows = [row for row in parsed.rows if row.cells]
        if not rows:
            continue
        occupied, placements, max_cols = set(), [], 0
        for row_idx, parsed_row in enumerate(rows):
            col_idx = 0
            for parsed_cell in parsed_row.cells:
                colspan = _safe_int(parsed_cell.attrs.get("colspan"), 1)
                rowspan = min(_safe_int(parsed_cell.attrs.get("rowspan"), 1), len(rows) - row_idx)
                while any((row_idx, col_idx + delta) in occupied for delta in range(colspan)):
                    col_idx += 1
                placements.append((row_idx, col_idx, rowspan, colspan, parsed_cell))
                for r in range(row_idx, row_idx + rowspan):
                    for c in range(col_idx, col_idx + colspan):
                        occupied.add((r, c))
                max_cols = max(max_cols, col_idx + colspan)
                col_idx += colspan
        # Estimate minimum readable widths from unbreakable words and Chinese text.
        widths = [0.9] * max_cols
        for _r, c, _rs, cs, cell in placements:
            words = re.findall(r'[A-Za-z0-9_.%+-]+|[^\x00-\x7f]', re.sub('<[^>]+>', '', cell.text))
            longest = max((len(word) for word in words), default=1)
            need = min(5.0, max(0.9, longest * .14 + .35)) / cs
            for column in range(c, c+cs):
                widths[column] = max(widths[column], need)
        landscape = max_cols >= 9 or sum(widths) > 16.6
        if landscape:
            section = doc.add_section(WD_SECTION_START.NEW_PAGE)
            section.orientation = WD_ORIENT.LANDSCAPE
            section.page_width, section.page_height = Cm(29.7), Cm(21)
        available = 25.3 if landscape else 16.6
        table = doc.add_table(rows=len(rows), cols=max_cols)
        table.autofit = False
        for column, width in zip(table.columns, widths):
            column.width = Cm(available * width / sum(widths))
        try:
            table.style = "Table Grid"
        except Exception:
            pass

        for row_idx, col_idx, rowspan, colspan, parsed_cell in placements:
            cell = table.cell(row_idx, col_idx)
            if colspan > 1 or rowspan > 1:
                cell = cell.merge(table.cell(row_idx + rowspan - 1, col_idx + colspan - 1))
            cell.text = ""
            cell.width = Cm(available * sum(widths[col_idx:col_idx+colspan]) / sum(widths))
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.space_after = Pt(0)
            _add_text_runs(paragraph, parsed_cell.text, size=TABLE_SIZE, bold=(parsed_cell.tag == "th"))
        for index, row in enumerate(rows):
            props = table.rows[index]._tr.get_or_add_trPr()
            if max((len(c.text) for c in row.cells), default=0) < 600:
                props.append(OxmlElement('w:cantSplit'))
            explicit = all(c.tag == 'th' for c in row.cells)
            inferred = (index == 0 and len(row.cells) >= 2 and all(not re.search(r'\d', c.text) for c in row.cells)
                        and len(rows) > 1 and any(re.search(r'\d', c.text) for c in rows[1].cells))
            if (explicit or inferred) and (index == 0 or table.rows[index-1]._tr.xpath('./w:trPr/w:tblHeader')):
                props.append(OxmlElement('w:tblHeader'))
        doc.add_paragraph()
        if landscape:
            section = doc.add_section(WD_SECTION_START.NEW_PAGE)
            section.orientation = WD_ORIENT.PORTRAIT
            section.page_width, section.page_height = Cm(21), Cm(29.7)
    return True


def _image_candidates(rel_path: str, images_search_root: Path):
    rel = Path(rel_path)
    if rel.is_absolute():
        yield rel
        return
    yield images_search_root / rel
    yield images_search_root.parent / rel
    yield images_search_root / "images" / rel.name


def _add_image(doc, image_line: str, images_search_root: Path | None):
    m = re.match(r"!\[[^\]]*\]\(([^)]+)\)", image_line)
    if not (m and images_search_root):
        return "图片路径无法识别。"
    rel_path = m.group(1)
    for ip in _image_candidates(rel_path, images_search_root):
        ip = ip.resolve()
        if ip.exists():
            try:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run()
                from PIL import Image
                with Image.open(ip) as image:
                    width, height = image.size
                section = doc.sections[-1]
                max_width = section.page_width - section.left_margin - section.right_margin
                max_width = min(max_width, Inches(width / 180))
                max_height = section.page_height - section.top_margin - section.bottom_margin - Cm(1.7)
                scale = min(max_width / width, max_height / height)
                r.add_picture(str(ip), width=int(width*scale), height=int(height*scale))
                p.paragraph_format.keep_together = True
                if height*scale > max_height*.85:
                    p.paragraph_format.page_break_before = True
            except Exception:
                return f"Word 无法嵌入图片：{rel_path}"
            return None
    return f"Word 找不到图片：{rel_path}"


def _configure_document(doc):
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    style.paragraph_format.line_spacing = LINE_SPACING
    style.paragraph_format.widow_control = True

    for lv in range(1, 5):
        try:
            hs = doc.styles[f"Heading {lv}"]
        except KeyError:
            continue
        hs.font.name = "Times New Roman"
        hs.font.size = Pt(HEADING_SIZES.get(lv, 12))
        hs.font.bold = True
        hs.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        hs.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        hs.paragraph_format.line_spacing = 1.3
        hs.paragraph_format.keep_with_next = True
        hs.paragraph_format.keep_together = True

    for sec in doc.sections:
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.top_margin = Cm(2.4)
        sec.bottom_margin = Cm(2.8)
        sec.left_margin = sec.right_margin = Cm(2.2)


def _field(paragraph, instruction, placeholder='请在 Word 中更新域'):
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), instruction)
    run = OxmlElement('w:r')
    text = OxmlElement('w:t')
    text.text = placeholder
    run.append(text)
    field.append(run)
    paragraph._p.append(field)


def _furniture(doc, layout):
    from PIL import Image
    mode = layout.get('header_mode', 'simple')
    if mode == 'off':
        return
    for section in doc.sections:
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False
        section.header_distance = section.footer_distance = Cm(.6)
        for kind, container in [('header', section.header), ('footer', section.footer)]:
            for old in list(container._element):
                container._element.remove(old)
            container.add_paragraph()
            p = container.paragraphs[0]
            p.clear()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            picture = layout.get('furniture', {}).get(kind)
            if picture and mode == 'original':
                with Image.open(picture) as im:
                    w, h = im.size
                scale = min((section.page_width-section.left_margin-section.right_margin)/w, Cm(1.3)/h)
                p.add_run().add_picture(picture, width=int(w*scale), height=int(h*scale))
            elif kind == 'header':
                set_font(p.add_run(layout.get('title', '译文')), size=8)
            if kind == 'footer':
                page = container.add_paragraph()
                page.alignment = WD_ALIGN_PARAGRAPH.CENTER
                page.paragraph_format.space_before = page.paragraph_format.space_after = Pt(0)
                _field(page, 'PAGE', '1')
    update = OxmlElement('w:updateFields')
    update.set(qn('w:val'), 'true')
    doc.settings.element.append(update)


def make_docx(md_text: str, docx_path, images_search_root=None, *, layout=None):
    doc = Document()
    _configure_document(doc)
    root = Path(images_search_root) if images_search_root else None

    lines = md_text.splitlines()
    i = 0
    reference_level = None
    warnings = []
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s:
            i += 1
            continue

        if "<table" in s.lower():
            block = [line]
            while "</table>" not in lines[i].lower() and i + 1 < len(lines):
                i += 1
                block.append(lines[i])
            if _add_html_table(doc, "\n".join(block)):
                i += 1
                continue

        images = list(re.finditer(r"!\[[^\]]*\]\([^)]+\)", s))
        if images:
            position = 0
            for image in images:
                before = s[position:image.start()].strip()
                if before:
                    _paragraph(doc, before)
                warning = _add_image(doc, image[0], root)
                if warning:
                    warnings.append(warning)
                after_lines = [v.strip() for v in lines[i+1:] if v.strip()]
                if not warning and after_lines and re.match(r'^(?:图|表|Figure|Table)\s*\d', after_lines[0], re.I):
                    doc.paragraphs[-1].paragraph_format.keep_with_next = True
                position = image.end()
            if s[position:].strip():
                _paragraph(doc, s[position:].strip())
            i += 1
            continue

        hm = re.match(r"^(#{1,6})\s+(.+)$", line)
        if hm:
            lv = len(hm.group(1))
            text = hm.group(2).strip()
            p = doc.add_heading(text, level=lv)
            for r in p.runs:
                set_font(r, bold=True)
            if _is_reference_heading(text):
                reference_level = lv
            elif reference_level is not None and lv <= reference_level:
                reference_level = None
            if text.strip().lower() in {'目录', 'table of contents', 'contents'}:
                end = i+1
                while end < len(lines) and not re.match(r'^#{1,6}\s+', lines[end]):
                    end += 1
                entries = [x.strip() for x in lines[i+1:end] if x.strip() and not x.strip().isdigit()]
                headings = {re.sub(r'^#{1,6}\s+', '', x).strip() for x in lines[end:] if re.match(r'^#{1,6}\s+', x)}
                if entries and all(x in headings for x in entries):
                    _field(doc.add_paragraph(), 'TOC \\o "1-3" \\h \\z \\u')
                    i = end
                    continue
                _paragraph(doc, '以下目录如含页码，属于原文页码，不代表译文分页。', size=9, first_indent=False)
            i += 1
            continue

        if reference_level is not None:
            _reference_paragraph(doc, s)
            i += 1
            continue

        if re.match(r"^(图|表)\s*\d", s):
            caption = _paragraph(doc, s, size=9, first_indent=False, italic=True)
            following = next((x.strip() for x in lines[i+1:] if x.strip()), '')
            previous = next((x.strip() for x in reversed(lines[:i]) if x.strip()), '')
            if (following.startswith('![') or following.lower().startswith('<table')) and not previous.startswith('!['):
                caption.paragraph_format.keep_with_next = True
            caption.paragraph_format.keep_together = True
            i += 1
            continue

        if re.match(r"^\d+\.\s+", s):
            p = _paragraph(doc, s, size=8, first_indent=False)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.0
            i += 1
            continue

        _paragraph(doc, s)
        i += 1

    _furniture(doc, layout or {'header_mode': 'simple'})
    warnings.extend((layout or {}).get('warnings', []))
    doc.save(str(docx_path))
    print(f"Wrote {docx_path}")
    return warnings


if __name__ == "__main__":
    args = sys.argv[1:]
    if not args:
        print("Usage: python make_docx.py <input.md> [output.docx] [images_root]")
        sys.exit(1)
    md_path = Path(args[0])
    docx_path = Path(args[1]) if len(args) > 1 else md_path.with_suffix(".docx")
    img_root = Path(args[2]) if len(args) > 2 else md_path.parent
    make_docx(md_path.read_text(encoding="utf-8"), docx_path, img_root)
