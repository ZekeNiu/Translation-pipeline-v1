import json
from pathlib import Path
import shutil
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from task_paths import artifact, translation_lock, PUBLIC
from task_state import write_json, read_json, fingerprint, file_hash
from review import reexport, confirm_unit, adopt_omission, all_units
from quality_report import active_issues
from mineru_sidecar import load_mineru_sidecar, strip_excluded_lines, find_omissions
from make_docx import make_docx
from pdf_helpers import make_pdf


class OutputTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.out = self.root / 'out'
        self.out.mkdir()

    def document(self):
        unit = {'id': 's:b0', 'source': 'Force 3.', 'translated': '力为 4。', 'kind': 'text',
                'chapter': 'Methods', 'page': 2, 'bbox': [100, 200, 800, 300], 'segment': 1,
                'source_hash': fingerprint('Force 3.'), 'issues': [], 'glossary_terms': []}
        pdf = make_pdf(self.root / 'origin.pdf', 3)
        document = {'version': 1, 'identity': 'test', 'source': {'path': str(pdf), 'hash': file_hash(pdf)},
                    'config': {'provider': 'custom', 'model': 'test', 'base_url': 'https://example.invalid', 'temperature': .1},
                    'segments': [{'id': 's', 'units': [unit], 'previous_context': '', 'next_context': ''}]}
        write_json(self.out / 'review_document.json', document)
        write_json(self.out / 'task_state.json', {'identity': 'test'})
        write_json(self.out / 'export_options.json', {'header_mode': 'simple'})
        return document, unit

    def test_migration_backup_resume_links_and_clean_root(self):
        self.document()
        (self.out / 'images').mkdir()
        (self.out / 'images' / 'x.png').write_bytes(b'image')
        (self.out / 'translated.md').write_text('![](images/x.png)', encoding='utf-8')
        with translation_lock(self.out):
            self.assertTrue(artifact(self.out, 'task_state.json').exists())
        self.assertEqual((self.out / 'translated.md').read_text(), '![](_internal/images/x.png)')
        self.assertEqual((self.out / '_internal/legacy_backup/translated.md').read_text(), '![](images/x.png)')
        with translation_lock(self.out):
            pass
        self.assertTrue({p.name for p in self.out.iterdir()} <= PUBLIC | {'_internal'})

    def test_interrupted_migration_recovers_without_overwriting_backup(self):
        self.document()
        import task_paths
        original = task_paths.os.replace
        def stop(source, target):
            if Path(source).name == 'task_state.json':
                raise OSError('interrupted')
            return original(source, target)
        with patch('task_paths.os.replace', side_effect=stop):
            with self.assertRaises(OSError), translation_lock(self.out):
                pass
        with translation_lock(self.out):
            self.assertTrue(artifact(self.out, 'review_document.json').exists())
        self.assertTrue((self.out / '_internal/legacy_backup/task_state.json').is_file())

    def test_html_report_is_located_escaped_and_offline_after_move(self):
        document, unit = self.document()
        unit['source'] += '<script>alert(1)</script>'
        write_json(self.out / 'review_document.json', document)
        with patch('translate.call_chat_completion', side_effect=AssertionError('offline')):
            reexport(self.out)
        html = (self.out / 'quality_report.html').read_text(encoding='utf-8')
        self.assertIn('原 PDF 第 2 页', html)
        self.assertIn('力为 4', html)
        self.assertNotIn('<script>alert', html)
        self.assertTrue(artifact(self.out, '_review_pages').joinpath('page_2.pdf').exists())
        moved = self.root / 'moved'
        shutil.copytree(self.out, moved)
        Path(document['source']['path']).unlink()
        with patch('translate.call_chat_completion', side_effect=AssertionError('offline')):
            reexport(moved)
        self.assertIn('打开原第 2 页', (moved / 'quality_report.html').read_text(encoding='utf-8'))

    def test_confirmed_issue_reappears_after_translation_change(self):
        document, unit = self.document()
        reexport(self.out)
        confirm_unit(self.out, unit['id'])
        edits = read_json(artifact(self.out, 'review_edits.json'))
        self.assertEqual(active_issues(unit, edits), [])
        unit['translated'] = '力为 5。'
        self.assertIn('数字可能发生变化', active_issues(unit, edits))

    def test_export_failure_keeps_previous_artifacts_and_can_retry(self):
        self.document()
        reexport(self.out)
        old = (self.out / 'translated.docx').read_bytes()
        with patch('make_docx.make_docx', side_effect=OSError('locked')):
            with self.assertRaises(OSError):
                reexport(self.out)
        self.assertEqual(old, (self.out / 'translated.docx').read_bytes())
        reexport(self.out)
        self.assertFalse(read_json(artifact(self.out, 'quality_report.json'))['review_export_pending'])

    def test_selected_omission_adoption_is_durable_and_idempotent(self):
        document, unit = self.document()
        omission = {'id': 'missing1', 'source': 'Force 2.', 'page': 2, 'bbox': [100,100,800,150],
                    'insert_before': unit['id'], 'reason': '解析疑似遗漏'}
        document['omissions'] = [omission]
        write_json(self.out / 'review_document.json', document)
        with patch('translate.call_chat_completion', side_effect=AssertionError('no model on adoption')):
            adopt_omission(self.out, 'missing1', '力为 2。')
            reexport(self.out)
        text = (self.out / 'translated.md').read_text(encoding='utf-8')
        self.assertEqual(text.count('力为 2'), 1)
        self.assertLess(text.index('力为 2'), text.index('力为 4'))
        with self.assertRaises(ValueError):
            adopt_omission(self.out, 'missing1', '力为 2。')

    def test_parallel_export_uses_same_lock(self):
        self.document()
        with translation_lock(self.out):
            with self.assertRaises(RuntimeError):
                reexport(self.out)

    def test_settings_and_export_share_lock_and_preserve_previous_settings(self):
        self.document()
        from task_state import task_lock
        from make_docx import make_docx as real_export
        def guarded(*args, **kwargs):
            with self.assertRaises(RuntimeError), task_lock(artifact(self.out, '_export_lock')):
                pass
            # Export owns a snapshot, allowing new manual edits during rendering.
            with task_lock(self.out / '_internal'):
                pass
            self.assertEqual(kwargs['layout']['header_mode'], 'off')
            return real_export(*args, **kwargs)
        with patch('make_docx.make_docx', side_effect=guarded):
            reexport(self.out, export_options={'header_mode': 'off'})
        backups = list(artifact(self.out, '_review_history').glob('*/export_options.json'))
        self.assertTrue(any(read_json(p)['header_mode'] == 'simple' for p in backups))

    def test_edit_during_export_is_saved_for_next_snapshot(self):
        self.document()
        from review import save_edit
        from review_state import revision_state
        from make_docx import make_docx as real_export
        def edit_during_render(*args, **kwargs):
            save_edit(self.out, 's:b0', '测得力为 3。', export=False)
            return real_export(*args, **kwargs)
        with patch('make_docx.make_docx', side_effect=edit_during_render):
            reexport(self.out)
        self.assertIn('力为 4', (self.out / 'translated.md').read_text(encoding='utf-8'))
        state = revision_state(self.out)
        self.assertGreater(state['revision'], state['exported_revision'])
        reexport(self.out)
        self.assertIn('测得力为 3', (self.out / 'translated.md').read_text(encoding='utf-8'))


class ParseAndLayoutTests(unittest.TestCase):
    def test_margin_evidence_preserves_body_with_same_text(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            records = [{'type':'header','text':'Repeated title','bbox':[100,20,800,50],'page_idx':i} for i in range(2)]
            records += [{'type':'text','text':'Repeated title','bbox':[100,200,800,250],'page_idx':1}]
            records += [{'type':'footer','text':'Journal Footer','bbox':[100,950,800,990],'page_idx':i} for i in range(2)]
            records += [{'type':'header','text':'A whole body paragraph that was misclassified.', 'bbox':[100,120,800,300],'page_idx':0}]
            write_json(root/'test_content_list.json',records)
            sidecar=load_mineru_sidecar(root)
            text, count=strip_excluded_lines('Repeated title\nJournal Footer\n',sidecar)
            self.assertIn('Repeated title',text)
            self.assertEqual(count,1)
            candidates=find_omissions(sidecar,text,[{'id':'next','page':1,'bbox':[100,350,800,500],'kind':'text'}])
            self.assertEqual(len(candidates),1)
            self.assertEqual(candidates[0]['insert_before'],'next')

    def test_a4_wide_table_fields_and_long_rows(self):
        with tempfile.TemporaryDirectory() as td:
            path=Path(td)/'out.docx'
            header='<tr>'+''.join('<th>MeasurementColumn</th>' for _ in range(13))+'</tr>'
            data='<tr>'+''.join('<td>12.3</td>' for _ in range(13))+'</tr>'
            long='<tr><td colspan="13">'+('long text '*100)+'</td></tr>'
            make_docx('# 标题\n\n<table>'+header+data+long+'</table>\n\n末段',path)
            doc=Document(path)
            self.assertAlmostEqual(doc.sections[0].page_width.cm,21,places=1)
            self.assertTrue(any(s.page_width > s.page_height for s in doc.sections))
            self.assertTrue(doc.element.xpath('//w:tblHeader'))
            self.assertTrue(doc.tables[0].rows[1]._tr.xpath('./w:trPr/w:cantSplit'))
            self.assertFalse(doc.tables[0].rows[2]._tr.xpath('./w:trPr/w:cantSplit'))
            for section in doc.sections:
                fields=section.footer._element.xpath('.//w:fldSimple')
                self.assertEqual(len(fields),1)

    def test_image_caption_pair_and_safe_dimensions(self):
        from PIL import Image
        with tempfile.TemporaryDirectory() as td:
            root=Path(td);Image.new('RGB',(2000,3000),'white').save(root/'figure.png')
            make_docx('![](figure.png)\n\n图 1. 图注\n\n正文',root/'out.docx',root)
            doc=Document(root/'out.docx')
            self.assertTrue(doc.paragraphs[0].paragraph_format.keep_with_next)
            self.assertLess(doc.inline_shapes[0].height.cm,23)
            self.assertLessEqual(doc.inline_shapes[0].width.cm,16.6)

    def test_preview_does_not_apply_preferences(self):
        from export_layout import preview_layout
        with tempfile.TemporaryDirectory() as td:
            root=Path(td)
            with translation_lock(root):
                write_json(artifact(root,'export_options.json'),{'header_mode':'simple'})
                path=preview_layout(root,{'source':{}},{'header_mode':'off'})
            self.assertTrue(path.exists())
            self.assertEqual(read_json(artifact(root,'export_options.json'))['header_mode'],'simple')

    def test_numeric_equivalence_does_not_hide_changed_values(self):
        from review import text_issues
        self.assertNotIn('数字可能发生变化',text_issues('Angle $1 0 . 7 ^ { \\circ }$.','角度10.7°。'))
        self.assertIn('数字可能发生变化',text_issues('Angle $1 0 . 7 ^ { \\circ }$.','角度10.8°。'))

    def test_merged_table_gets_range_without_guessing_cell_pages(self):
        from review import make_document
        from mineru_sidecar import MinerUSidecar
        from translate import resolve_provider_config
        a, b = '<table><tr><td>One</td></tr></table>', '<table><tr><td>Two</td></tr></table>'
        merged = '<table><tr><td>One</td></tr><tr><td>Two</td></tr></table>'
        key = 's00000_' + fingerprint(merged)[:12]
        state = {'identity': 'x', 'segments': {key: {'alignment': [{'kind':'table','source':merged,'translated':merged}]}}}
        sidecar = MinerUSidecar(provenance=[{'text':a,'page':2,'bbox':None,'block_id':0}, {'text':b,'page':3,'bbox':None,'block_id':1}])
        doc = make_document(state,[merged],[merged],['Methods'],sidecar,{},resolve_provider_config())
        table = doc['segments'][0]['units'][0]
        self.assertEqual(table['pages'],[2,3])
        self.assertEqual(table['cells'][0]['pages'],[2,3])
        self.assertIsNone(table['cells'][0]['bbox'])

    def test_duplicate_formula_font_groups_render_without_latex_commands(self):
        from inline_semantics import normalize_inline_output
        self.assertEqual(normalize_inline_output(r'$5 ^ { \mathrm { { t h } } }$'), '5 ᵗʰ')
