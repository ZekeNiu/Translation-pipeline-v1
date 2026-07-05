import json
import re
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document

from latex_utils import latex_to_readable_text
from make_docx import make_docx
from mineru_sidecar import load_mineru_sidecar, strip_excluded_lines
import translate
from translate import ProviderConfig


class PipelineQualityTests(unittest.TestCase):
    def test_scientific_comparators_are_not_stripped(self):
        text = "HIT includes short (<45 s) to long (>2 min) intervals."
        self.assertEqual(translate.normalize_source_text(text), text)

    def test_abstract_heading_splits_inline_body(self):
        text = "Abstract High-intensity interval training is useful.\n\n## 1 Intro"
        out = translate.normalize_abstract_heading(text)
        self.assertIn("## 摘要\n\nHigh-intensity interval training is useful.", out)
        self.assertNotIn("## 摘要 High-intensity", out)

    def test_math_placeholders_restore_and_clean_latex_is_safe(self):
        text = r"$v_{VO2max}$ and $\frac{a}{b}$ and $T@\dot{V}O_2max$"
        protected, placeholders = translate.protect_fragments(text)
        self.assertNotIn(r"\frac{a}{b}", protected)
        restored = translate.restore_placeholders(protected, placeholders)
        self.assertIn("a/b", restored)
        self.assertIn("T@V̇O₂max", restored)
        self.assertNotIn(r"\frac", restored)
        self.assertEqual(translate.clean_latex(text), text)

    def test_reference_headings_are_isolated(self):
        headings = ["## References", "# References", "## REFERENCES", "## Bibliography", "## 参考文献"]
        for heading in headings:
            with self.subTest(heading=heading):
                body, refs = translate.split_references(f"## Body\ntext\n\n{heading}\n1. Smith")
                self.assertIn("## Body", body)
                self.assertTrue(refs.startswith(heading))
                chunks, refs2 = translate.segment_document(f"## Body\ntext\n\n{heading}\n1. Smith")
                self.assertEqual(len(chunks), 1)
                self.assertTrue(refs2.startswith(heading))

    def test_reference_continuation_lines_are_merged(self):
        refs = (
            "## References\n\n"
            "15. Krustrup P, Ortenblad N. Maximal voluntary contraction force during\n\n"
            "the first 72 h after a high-level competitive soccer game. Eur J Appl Physiol. 2011;111:2987-95.\n\n"
            "[16] Next Author. Next title. Journal. 2000;1:1-2.\n"
        )
        normalized, report = translate.normalize_references(refs)
        self.assertIn(
            "15. Krustrup P, Ortenblad N. Maximal voluntary contraction force during the first 72 h",
            normalized,
        )
        self.assertIn("[16] Next Author.", normalized)
        self.assertNotIn("\n\nthe first 72 h", normalized)
        self.assertEqual(report["entries"], 2)
        self.assertEqual(report["continuation_lines_merged"], 1)

    def test_provider_endpoint_appends_chat_completions(self):
        cfg = translate.resolve_provider_config(
            provider="custom",
            base_url="https://example.test/v1",
            api_key="test-key",
            model="test-model",
        )
        self.assertEqual(cfg.endpoint(), "https://example.test/v1/chat/completions")

    def test_docx_renders_html_table_as_word_table(self):
        md = "# Title\n\n<table><tr><th>Mode</th><th>Duration</th></tr><tr><td>HIT</td><td>&lt;45 s</td></tr></table>"
        with tempfile.TemporaryDirectory() as td:
            out = Path(td) / "out.docx"
            make_docx(md, out, Path(td))
            doc = Document(str(out))
            self.assertEqual(len(doc.tables), 1)
            self.assertEqual(doc.tables[0].cell(1, 1).text, "<45 s")

    def test_latex_readable_normalizer_handles_mineru_formula(self):
        source = r"$[ \mathrm { v / p } \dot { V } \mathrm { O } _ { 2 \operatorname* { m a x } } ]$"
        self.assertEqual(latex_to_readable_text(source), "[v/p V̇O₂max]")
        self.assertEqual(latex_to_readable_text(r"$\frac{a}{b}$"), "a/b")
        self.assertEqual(latex_to_readable_text(r"$\geq 95$"), "≥ 95")
        self.assertEqual(latex_to_readable_text(r"$T@\dot{V}O_2max$"), "T@V̇O₂max")
        self.assertEqual(latex_to_readable_text(r"$V_{IFT}$"), "VIFT")

    def test_latex_text_macros_do_not_leak_as_words(self):
        source = r"$2 0 \mathrm { ~ s ~ } \textsuperscript { \textregistered } \textsuperscript { 1 7 0 \% }$"
        out = latex_to_readable_text(source)
        self.assertIn("®", out)
        self.assertIn("¹⁷⁰%", out)
        self.assertNotIn("textsuperscript", out)
        self.assertNotIn("textregistered", out)
        self.assertNotIn("mathrm", out)

    def test_docx_renders_formula_readably(self):
        formula = r"$[ \mathrm { v / p } \dot { V } \mathrm { O } _ { 2 \operatorname* { m a x } } ]$"
        with tempfile.TemporaryDirectory() as td:
            out = Path(td) / "formula.docx"
            make_docx(f"# Title\n\nThe metric is {formula}.", out, Path(td))
            text = "\n".join(p.text for p in Document(str(out)).paragraphs)
            self.assertIn("[v/p V̇O₂max]", text)
            self.assertNotIn(r"\mathrm", text)

    def test_mineru_sidecar_filters_header_footer_page_lines(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td)
            (root / "sample_content_list_v2.json").write_text(
                json.dumps(
                    [
                        [
                            {"type": "page_header", "content": {"paragraph_content": [{"type": "text", "content": "Journal Header"}]}},
                            {"type": "page_number", "text": "927"},
                            {"type": "equation_inline", "content": r"\dot { V } O_2"},
                        ]
                    ]
                ),
                encoding="utf-8",
            )
            sidecar = load_mineru_sidecar(root)
            cleaned, removed = strip_excluded_lines("Journal Header\n\nReal text\n927\n", sidecar)
            self.assertEqual(removed, 2)
            self.assertIn("Real text", cleaned)
            self.assertEqual(sidecar.formula_texts, [r"\dot { V } O_2"])

    def test_segment_translation_uses_cache(self):
        cfg = ProviderConfig(
            provider_name="custom",
            base_url="https://example.test/v1",
            api_key_env="AI_API_KEY",
            api_key="test-key",
            model="test-model",
        )
        calls = []

        def fake_call(_config, messages, timeout=None):
            calls.append(messages)
            match = re.search(r"<SOURCE>\n(.*?)\n</SOURCE>", messages[-1]["content"], re.DOTALL)
            return match.group(1) if match else messages[-1]["content"]

        with tempfile.TemporaryDirectory() as td, patch("translate.call_chat_completion", fake_call):
            cache_dir = Path(td)
            first = translate.translate_segment("A formula $x_1$ stays.", "", 1, 1, cfg, cache_dir)
            second = translate.translate_segment("A formula $x_1$ stays.", "", 1, 1, cfg, cache_dir)
            self.assertEqual(first, second)
            self.assertIn("x₁", first)
            self.assertEqual(len(calls), 1)

    def test_consistency_guide_does_not_preserve_body_phrases_in_english(self):
        guide = translate.build_consistency_guide(
            "## 2.1 Anaerobic Glycolytic Energy Contribution to HIT\n\n"
            "While the accumulated oxygen deficit is useful, high-intensity interval training matters."
        )
        self.assertIn("Reuse the same Chinese wording", guide)
        self.assertNotIn("Preserve these names/entities", guide)
        self.assertNotIn("Anaerobic Glycolytic Energy Contribution", guide)
        self.assertIn("high-intensity interval", "; ".join(translate.extract_term_audit_candidates(guide + " high-intensity interval training")))

    def test_untranslated_english_detection_finds_long_body_phrase(self):
        warnings = translate.find_untranslated_english("在本节中，Anaerobic Glycolytic Energy Contribution 对HIT参数具有依赖性。")
        self.assertIn("Anaerobic Glycolytic Energy Contribution", warnings)

    def test_table_cell_translation_preserves_structure(self):
        cfg = ProviderConfig(
            provider_name="custom",
            base_url="https://example.test/v1",
            api_key_env="AI_API_KEY",
            api_key="test-key",
            model="test-model",
        )

        def fake_call(_config, messages, timeout=None):
            items = json.loads(messages[-1]["content"])
            return json.dumps([item.replace("Duration", "时长").replace("Mode", "形式") for item in items])

        html = "<table><tr><td>Mode</td><td>Duration</td></tr></table>"
        with tempfile.TemporaryDirectory() as td, patch("translate.call_chat_completion", fake_call):
            out = translate.translate_tables_in_text(html, cfg, Path(td), lambda *_args, **_kw: None)
            self.assertIn("<table", out)
            self.assertIn("形式", out)
            self.assertIn("时长", out)

    def test_mixed_math_text_and_html_superscripts_remain_translatable(self):
        cfg = ProviderConfig(
            provider_name="custom",
            base_url="https://example.test/v1",
            api_key_env="AI_API_KEY",
            api_key="test-key",
            model="test-model",
        )
        seen_items = []

        def fake_call(_config, messages, timeout=None):
            items = json.loads(messages[-1]["content"])
            seen_items.extend(items)
            translated = []
            for item in items:
                translated.append(
                    item.replace("Work intensity", "工作强度")
                    .replace("Ground surface", "地面类型")
                    .replace("Sport specific", "专项运动")
                )
            return json.dumps(translated, ensure_ascii=False)

        html = (
            "<table><tr>"
            "<td>$Work intensity^a$</td>"
            "<td>Ground surface<sup>b</sup></td>"
            "<td>$Sport specific^f$</td>"
            "</tr></table>"
        )
        reports = []
        with tempfile.TemporaryDirectory() as td, patch("translate.call_chat_completion", fake_call):
            out = translate.translate_tables_in_text(
                html,
                cfg,
                Path(td),
                lambda *_args, **_kw: None,
                table_reports=reports,
            )
            self.assertTrue(any("Work intensity" in item for item in seen_items))
            self.assertIn("工作强度ᵃ", out)
            self.assertIn("地面类型ᵇ", out)
            self.assertIn("专项运动ᶠ", out)
            self.assertNotIn("<sup>", out)
            self.assertNotIn("$Work intensity", out)
            self.assertEqual(reports[0]["residual_english"], [])
            self.assertEqual(reports[0]["inline_artifacts"], [])

    def test_docx_uses_real_script_runs_for_unicode_scripts(self):
        with tempfile.TemporaryDirectory() as td:
            out = Path(td) / "scripts.docx"
            make_docx("# Title\n\n工作强度ᵃ and O₂.", out, Path(td))
            doc = Document(str(out))
            runs = doc.paragraphs[1].runs
            self.assertTrue(any(run.text == "ᵃ" and run.font.superscript for run in runs))
            self.assertTrue(any(run.text == "₂" and run.font.subscript for run in runs))

    def test_docx_reference_mode_uses_reference_style_for_continuations(self):
        with tempfile.TemporaryDirectory() as td:
            out = Path(td) / "refs.docx"
            make_docx("# Title\n\n## References\n\nContinuation without number.", out, Path(td))
            doc = Document(str(out))
            paragraph = doc.paragraphs[2]
            self.assertEqual(round(paragraph.runs[0].font.size.pt), 8)
            self.assertLess(paragraph.paragraph_format.first_line_indent, 0)

    def test_table_translation_falls_back_per_cell_not_whole_table(self):
        cfg = ProviderConfig(
            provider_name="custom",
            base_url="https://example.test/v1",
            api_key_env="AI_API_KEY",
            api_key="test-key",
            model="test-model",
        )
        calls = []

        def fake_call(_config, messages, timeout=None):
            items = json.loads(messages[-1]["content"])
            calls.append(items)
            if len(items) > 1:
                raise ValueError("batch too large")
            return json.dumps([items[0].replace("Format", "形式").replace("Work duration", "运动时长")])

        html = "<table><tr><td>Format</td><td>Work duration</td></tr></table>"
        reports = []
        with tempfile.TemporaryDirectory() as td, patch("translate.call_chat_completion", fake_call):
            out = translate.translate_tables_in_text(
                html,
                cfg,
                Path(td),
                lambda *_args, **_kw: None,
                table_reports=reports,
            )
            self.assertIn("形式", out)
            self.assertIn("运动时长", out)
            self.assertGreater(len(calls), 1)
            self.assertEqual(reports[0]["failures"], [])

    def test_chunking_merges_small_segments(self):
        body = "# Title\n\nshort\n\n## 1 Intro\n\n" + ("a" * 500) + "\n\n## 2 Methods\n\n" + ("b" * 500)
        chunks = translate.split_body_into_segments(body, max_chunk=10000, min_chunk=1800)
        self.assertLess(len(chunks), 3)

    def test_pipeline_prefers_full_md_and_writes_sidecar_summary(self):
        cfg_args = {
            "provider": "custom",
            "base_url": "https://example.test/v1",
            "model": "test-model",
            "api_key": "test-key",
        }

        def fake_call(_config, messages, timeout=None):
            content = messages[-1]["content"]
            if content.strip().startswith("["):
                return content
            match = re.search(r"<SOURCE>\n(.*?)\n</SOURCE>", content, re.DOTALL)
            return match.group(1) if match else content

        with tempfile.TemporaryDirectory() as td, patch("translate.call_chat_completion", fake_call):
            root = Path(td) / "mineru"
            root.mkdir()
            (root / "a.md").write_text("# Wrong\n\nWrong file", encoding="utf-8")
            (root / "full.md").write_text(
                "# Right\n\nJournal Header\n\nRight file\n\n"
                "## References\n\n1. Smith title during\n\ncontinuation line.\n\n2. Jones title.",
                encoding="utf-8",
            )
            (root / "doc_content_list_v2.json").write_text(
                json.dumps([[{"type": "page_header", "content": {"paragraph_content": [{"type": "text", "content": "Journal Header"}]}}]]),
                encoding="utf-8",
            )
            out = Path(td) / "out"
            md, _docx = translate.run(root, out, **cfg_args)
            result = md.read_text(encoding="utf-8")
            self.assertIn("# Right", result)
            self.assertNotIn("Wrong file", result)
            self.assertNotIn("Journal Header", result)
            self.assertIn("1. Smith title during continuation line.", result)
            self.assertTrue((out / "references_original.md").exists())
            self.assertTrue((out / "references_normalized.md").exists())
            self.assertIn("Reference normalization:", (out / "translation.log").read_text(encoding="utf-8"))
            self.assertTrue((out / "mineru_structure_summary.json").exists())


if __name__ == "__main__":
    unittest.main()
